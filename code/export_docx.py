# -*- coding: utf-8 -*-
"""用 python-docx 生成内嵌图片的 DOCX（替代 textutil —— 后者不内嵌图片）。

读取 paper/main.md 的 Markdown 子集，按结构写入 DOCX，并把 `![alt](src)` 解析为
真实的嵌入式图片（这是 T0-2 修复的核心）。
"""
from __future__ import annotations
import os, re, sys
from docx import Document
from docx.shared import Pt, Cm, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, os.path.join(ROOT, 'code'))
from export_paper import latex_to_unicode            # noqa: E402

PAPER = os.path.join(ROOT, 'paper', 'main.md')
FIGDIR = os.path.join(ROOT, 'paper', 'figs')
OUTDIR = os.path.join(ROOT, 'paper', 'export')
OUT = os.path.join(OUTDIR, '山区洪涝灾害下无人机运输与通信协同优化.docx')


def set_cjk(run):
    run.font.name = 'Songti SC'
    run._element.rPr.rFonts.set(qn('w:eastAsia'), '宋体')


def inline_runs(par, text):
    """把 **粗体**、`代码`、$公式$ 写入段落。"""
    text = re.sub(r'\$([^$]+)\$', lambda m: latex_to_unicode(m.group(1)), text)
    for tok in re.split(r'(\*\*[^*]+\*\*|`[^`]+`)', text):
        if not tok:
            continue
        if tok.startswith('**') and tok.endswith('**'):
            r = par.add_run(tok[2:-2]); r.bold = True
        elif tok.startswith('`') and tok.endswith('`'):
            r = par.add_run(tok[1:-1]); r.font.name = 'Menlo'; r.font.size = Pt(9)
        else:
            r = par.add_run(tok)
        set_cjk(r)


def add_table(doc, header, rows):
    t = doc.add_table(rows=1, cols=len(header))
    t.style = 'Table Grid'
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for j, h in enumerate(header):
        c = t.rows[0].cells[j]; c.text = ''
        inline_runs(c.paragraphs[0], h)
        for r in c.paragraphs[0].runs:
            r.bold = True
    for row in rows:
        cells = t.add_row().cells
        for j, v in enumerate(row[:len(header)]):
            cells[j].text = ''
            inline_runs(cells[j].paragraphs[0], v)
    doc.add_paragraph()


def build():
    md = open(PAPER, encoding='utf-8').read()
    doc = Document()
    # 页边距
    for sec in doc.sections:
        sec.top_margin = Cm(2.0); sec.bottom_margin = Cm(2.0)
        sec.left_margin = Cm(2.0); sec.right_margin = Cm(2.0)
    st = doc.styles['Normal']
    st.font.size = Pt(10.5)
    st.font.name = 'Songti SC'
    st.element.rPr.rFonts.set(qn('w:eastAsia'), '宋体')

    lines = md.split('\n')
    i = 0
    nimg = 0
    while i < len(lines):
        ln = lines[i]
        # 图片
        m = re.match(r'^!\[([^\]]*)\]\(([^)]+)\)\s*$', ln.strip())
        if m:
            src = m.group(2)
            path = os.path.join(ROOT, 'paper', src) if not os.path.isabs(src) else src
            if os.path.exists(path):
                p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p.add_run().add_picture(path, width=Inches(6.0))
                nimg += 1
                cap = doc.add_paragraph(); cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
                r = cap.add_run(m.group(1)); r.font.size = Pt(9); set_cjk(r)
            i += 1; continue
        # 表格
        if ln.strip().startswith('|') and i + 1 < len(lines) and \
           re.match(r'^\s*\|[\s:|-]+\|\s*$', lines[i + 1]):
            hdr = [c.strip() for c in ln.strip().strip('|').split('|')]
            i += 2
            rows = []
            while i < len(lines) and lines[i].strip().startswith('|'):
                rows.append([c.strip() for c in lines[i].strip().strip('|').split('|')])
                i += 1
            add_table(doc, hdr, rows)
            continue
        # 标题
        m = re.match(r'^(#{1,4})\s+(.*)$', ln)
        if m:
            lvl = len(m.group(1))
            h = doc.add_heading(level=lvl)
            inline_runs(h, m.group(2))
            i += 1; continue
        # 引用
        if ln.strip().startswith('>'):
            buf = []
            while i < len(lines) and lines[i].strip().startswith('>'):
                buf.append(lines[i].strip().lstrip('>').strip()); i += 1
            p = doc.add_paragraph(); p.paragraph_format.left_indent = Cm(0.6)
            inline_runs(p, ' '.join(buf))
            for r in p.runs:
                r.italic = True
            continue
        # 列表
        if re.match(r'^\s*([-*+]|\d+\.)\s+', ln):
            while i < len(lines) and re.match(r'^\s*([-*+]|\d+\.)\s+', lines[i]):
                item = re.sub(r'^\s*([-*+]|\d+\.)\s+', '', lines[i])
                p = doc.add_paragraph(style='List Bullet')
                inline_runs(p, item)
                i += 1
            continue
        if not ln.strip():
            i += 1; continue
        # 段落
        buf = [ln]; i += 1
        while i < len(lines) and lines[i].strip() and \
                not re.match(r'^(#{1,4}\s|\s*([-*+]|\d+\.)\s|\s*>|\s*\||```|!\[)', lines[i]) \
                and not re.match(r'^\s*(---|\*\*\*|___)\s*$', lines[i]):
            buf.append(lines[i]); i += 1
        p = doc.add_paragraph()
        p.paragraph_format.first_line_indent = Cm(0.74)
        inline_runs(p, ' '.join(buf))
        # 引用标记 [n] 上标化（简化：不改）
    os.makedirs(OUTDIR, exist_ok=True)
    doc.save(OUT)
    print("DOCX ->", OUT)
    print("  内嵌图片: %d 张" % nimg)
    return nimg


if __name__ == '__main__':
    build()
